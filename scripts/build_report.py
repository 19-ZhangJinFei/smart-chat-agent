"""用Codex文档依赖运行；个人信息由忽略目录JSON注入，源码不含身份信息。

学校原始DOCX未取得时仅生成格式与可见结构对应的初稿，不声称模板逐项保真。
"""
import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]


def explicit_fonts(properties, chinese, western):
    """清除默认主题字体，避免Word/渲染器覆盖学校指定字体。"""
    fonts = properties.get_or_add_rFonts()
    for attribute in ["asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"]:
        fonts.attrib.pop(qn(f"w:{attribute}"), None)
    for attribute, value in [("ascii", western), ("hAnsi", western), ("eastAsia", chinese)]:
        fonts.set(qn(f"w:{attribute}"), value)


def font(run, chinese="宋体", western="Times New Roman", size=12, bold=False):
    run.font.name = western
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    explicit_fonts(run._element.get_or_add_rPr(), chinese, western)


def field(paragraph, instruction, cached="1"):
    start, code, separate, end = [OxmlElement(tag) for tag in ["w:fldChar", "w:instrText", "w:fldChar", "w:fldChar"]]
    start.set(qn("w:fldCharType"), "begin")
    code.set(qn("xml:space"), "preserve")
    code.text = instruction
    separate.set(qn("w:fldCharType"), "separate")
    end.set(qn("w:fldCharType"), "end")
    for element in [start, code, separate]:
        paragraph.add_run()._r.append(element)
    font(paragraph.add_run(cached), size=9)
    paragraph.add_run()._r.append(end)


def architecture(path):
    image = Image.new("RGB", (1600, 750), "white")
    draw = ImageDraw.Draw(image)
    face = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 30)
    bold = ImageFont.truetype("C:/Windows/Fonts/arialbd.ttf", 28)
    boxes = [(60, 230, 350, 380, "Vue 3 browser", "3 modes / SSE"),
             (480, 230, 770, 380, "Nginx", "static + /api"),
             (900, 230, 1200, 380, "FastAPI", "lock / route / save"),
             (1260, 40, 1560, 180, "Qwen + tools", "create_agent"),
             (900, 510, 1200, 660, "Business SQLite", "pairs + version"),
             (1260, 510, 1560, 660, "Checkpoint SQLite", "thread state")]
    for x1, y1, x2, y2, title, detail in boxes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=12, fill="#F3F6FC", outline="#263D65", width=3)
        draw.text(((x1+x2)/2, y1+45), title, fill="#172D51", font=bold, anchor="mm")
        draw.text(((x1+x2)/2, y1+96), detail, fill="#334155", font=face, anchor="mm")
    for a, b in [((350,305),(480,305)), ((770,305),(900,305)), ((1200,270),(1260,140)),
                 ((1050,380),(1050,510)), ((1200,355),(1350,510))]:
        draw.line((a,b), fill="#334155", width=5)
        import math
        angle = math.atan2(b[1]-a[1], b[0]-a[0])
        draw.polygon([b, (b[0]-20*math.cos(angle-0.4), b[1]-20*math.sin(angle-0.4)),
                      (b[0]-20*math.cos(angle+0.4), b[1]-20*math.sin(angle+0.4))], fill="#334155")
    image.save(path)


def build(profile, output):
    info = json.loads(profile.read_text(encoding="utf-8-sig"))
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(2.54)
    section.header_distance = section.footer_distance = Cm(1.27)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(12)
    explicit_fonts(normal._element.get_or_add_rPr(), "宋体", "Times New Roman")
    fmt = normal.paragraph_format
    fmt.line_spacing, fmt.space_after, fmt.first_line_indent = 1.5, Pt(6), Pt(24)
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # 2字符缩进使用OOXML字符值，避免用空格。
    normal._element.get_or_add_pPr().find(qn("w:ind")).set(qn("w:firstLineChars"), "200")
    for level, size, before, after in [(1,16,18,12),(2,14,12,8),(3,12,10,6)]:
        style = document.styles[f"Heading {level}"]
        style.font.name, style.font.size, style.font.bold = "Arial", Pt(size), True
        style.font.color.rgb = RGBColor(0,0,0)
        explicit_fonts(style._element.get_or_add_rPr(), "黑体", "Arial")
        explicit_fonts(document.styles[f"Heading {level} Char"]._element.get_or_add_rPr(), "黑体", "Arial")
        sf = style.paragraph_format
        sf.first_line_indent, sf.line_spacing = Pt(0), 1.0
        sf.space_before, sf.space_after, sf.keep_with_next = Pt(before), Pt(after), True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.paragraph_format.first_line_indent = Pt(0)
    font(header.add_run("《企业级软件实践与工程能力提升训练营》项目报告"), size=9)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k,v in {"val":"single","sz":"4","color":"000000","space":"4"}.items():
        bottom.set(qn(f"w:{k}"), v)
    border.append(bottom)
    header._p.get_or_add_pPr().append(border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Pt(0)
    font(footer.add_run("第 "), size=9); field(footer," PAGE ")
    font(footer.add_run(" 页 共 "), size=9); field(footer," NUMPAGES ","15")
    font(footer.add_run(" 页"), size=9)
    settings = document.settings.element
    update = OxmlElement("w:updateFields"); update.set(qn("w:val"),"true"); settings.append(update)
    document.core_properties.title = "智聊 基于LangChain的智能客服系统项目报告"
    document.core_properties.author = info["姓名"]
    document.core_properties.comments = "报告初稿，需本人复核与原始学校DOCX模板核对"

    def p(text):
        return document.add_paragraph(text)
    def h(text, level=2):
        return document.add_heading(text,level)
    def caption(text):
        para = p(text); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.first_line_indent = Pt(0)
        para.paragraph_format.line_spacing = 1.0
        for r in para.runs: font(r,size=10.5,bold=True)
        return para
    def table(title, heads, rows, widths):
        cap = caption(title); cap.paragraph_format.keep_with_next = True
        t = document.add_table(rows=1, cols=len(heads)); t.autofit=False
        for i,width in enumerate(widths): t.columns[i].width=Cm(width)
        for i,text in enumerate(heads): t.rows[0].cells[i].text=text
        for row in rows:
            cells=t.add_row().cells
            for i,text in enumerate(row): cells[i].text=str(text)
        borders=OxmlElement("w:tblBorders")
        for edge in ["top","left","bottom","right","insideH","insideV"]:
            e=OxmlElement(f"w:{edge}")
            for k,v in {"val":"single","sz":"4","color":"D9D9D9"}.items():e.set(qn(f"w:{k}"),v)
            borders.append(e)
        t._tbl.tblPr.append(borders)
        for j,row in enumerate(t.rows):
            props=row._tr.get_or_add_trPr(); props.append(OxmlElement("w:cantSplit"))
            if j==0:props.append(OxmlElement("w:tblHeader"))
            for i,cell in enumerate(row.cells):
                cell.width=Cm(widths[i])
                margins=OxmlElement("w:tcMar")
                for edge in ["top","left","bottom","right"]:
                    item=OxmlElement(f"w:{edge}");item.set(qn("w:w"),"65");item.set(qn("w:type"),"dxa");margins.append(item)
                cell._tc.get_or_add_tcPr().append(margins)
                if j==0:
                    shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"DDEBF7");cell._tc.get_or_add_tcPr().append(shade)
                for para in cell.paragraphs:
                    para.paragraph_format.first_line_indent=Pt(0);para.paragraph_format.line_spacing=1.0
                    para.paragraph_format.space_after=Pt(2)
                    para.alignment=WD_ALIGN_PARAGRAPH.CENTER if j==0 or i==0 else WD_ALIGN_PARAGRAPH.LEFT
                    for r in para.runs:font(r,size=10.5,bold=j==0)
        return t
    def picture(file, title, height=None):
        para=p("");para.alignment=WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.first_line_indent=Pt(0);para.paragraph_format.keep_with_next=True
        para.add_run().add_picture(str(file),width=Cm(15.6))
        caption(title)
    def code(filename, content):
        para=p("关键代码："+filename);para.paragraph_format.first_line_indent=Pt(0)
        for r in para.runs:font(r,bold=True)
        para=p(content);para.paragraph_format.first_line_indent=Pt(0)
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.line_spacing=1.0;para.paragraph_format.space_after=Pt(6)
        para.paragraph_format.keep_together=True
        for r in para.runs:font(r,western="Consolas",size=10.5)

    # 学校封面，与正文严格分开。
    for _ in range(3): p("")
    para=p("《企业级软件实践与工程能力提升训练营》");para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.first_line_indent=Pt(0)
    font(para.runs[0],chinese="黑体",size=18,bold=True)
    para=p("项目报告");para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.first_line_indent=Pt(0)
    font(para.runs[0],chinese="黑体",size=36,bold=True)
    for _ in range(2):p("")
    cover = {"系统名称":"智聊——基于LangChain的智能客服系统", **info}
    for label,value in cover.items():
        para=p("");para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.first_line_indent=Pt(0)
        para.paragraph_format.space_after=Pt(12)
        font(para.add_run(label+"　"),chinese="黑体",size=14,bold=True)
        run=para.add_run(value);font(run,size=13);run.underline=True
    document.add_page_break()
    para=p("目录");para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.first_line_indent=Pt(0)
    font(para.runs[0],chinese="黑体",size=16,bold=True)
    toc=p("");toc.paragraph_format.first_line_indent=Pt(0)
    field(toc,' TOC \\o "1-2" \\h \\z \\u ',"")
    document.add_page_break()

    h("1　需求分析",1)
    h("1.1　项目背景与意义")
    p("智聊面向电商客服教学场景，将自然语言对话与可验证的业务工具结合。单纯语言模型可以组织回答，却可能编造订单状态、天气或计算结果。本项目通过工具查询固定样例，在回答中声明模拟业务边界，并以数据库保存多轮上下文。目标是形成可运行、可现场演示、可复现测试的个人课设，而不是只展示接口截图。")
    h("1.2　用户与业务场景")
    p("使用者通过浏览器提出订单、天气、计算、时间和优惠券问题，也可普通闲聊。系统提供自动、普通和智能体三模式，既适合连续使用，也便于课堂对比模型调用方式。会话可新建、改名、删除，刷新和服务重启后继续读取历史。项目暂不引入账号体系，部署只监听本机，适用于单人演示。")
    h("1.3　功能需求")
    table("表1　功能需求清单",["模块","需求","验收方式"],[
        ("会话","增删查改、持久化、隔离","接口与重启测试"),("对话","多轮记忆与三模式","真实模型追问"),
        ("工具","订单、天气、计算、时间、优惠券","固定题评估"),("体验","普通SSE、标签与重试","浏览器与跨块测试"),
        ("交付","容器、备份与恢复","停启和独立目录恢复")],[2,7.5,6.1])
    h("1.4　非功能需求")
    p("系统要求输入校验、友好错误、同会话生成互斥和每IP滑动窗口限流。失败轮次不写成成功问答，真实Key不进入公开仓库、镜像或浏览器。过程材料只记录真实结果，离线回归与真实模型评估分开说明。")

    document.add_page_break()
    h("2　系统设计",1);h("2.1　总体架构设计")
    p("系统总体架构如图1所示。Vue负责交互与SSE显示，Nginx托管构建产物并代理API，FastAPI协调会话、意图和模型。普通调用经LangChain模型接口；智能体经create_agent调用工具。业务SQLite与检查点SQLite共同挂载到宿主机数据目录。")
    diagram=output.parent/"architecture.png";architecture(diagram)
    picture(diagram,"图1　系统总体架构图")
    p("请求先验证会话和输入并取得会话锁。模型成功返回最终正文后，后端在一个事务中保存问题与答案，同时递增历史版本和活跃时间。工具元数据独立保存用于界面展示。自动模式先分类，再调用普通或智能体路径，分类错误会明确返回，不能假装已成功路由。")
    h("2.2　技术选型与理由")
    table("表2　技术栈与职责",["技术","用途"],[
        ("Python3.11 / FastAPI","分层API、校验与应用生命周期"),("SQLAlchemy / SQLite","业务事务和本地持久化"),
        ("LangChain / LangGraph","模型集成、工具循环与检查点"),("Vue3 / Element Plus / Vite","会话界面、组件和构建"),
        ("Docker Compose / Nginx","本机交付、代理和数据挂载")],[6,9.6])

    document.add_page_break()
    h("2.3　数据结构设计")
    p("业务库存储会话和消息。会话ID是稳定标识，history_version只在成功轮次递增；manual_title保护人工标题。消息保存角色、正文、时间以及路由和工具名，按ID排序，避免同一时间戳产生展示顺序不确定。删除会话通过外键级联清理消息，并使用检查点库官方delete_thread清理线程。")
    table("表3　主要数据字段",["实体","关键字段","作用"],[
        ("会话","id、title、manual_title","标识、标题与人工改名"),("会话","history_version、updated_at","恢复判断和活跃排序"),
        ("消息","session_id、role、content","外键、角色与正文"),("消息","route、tools_used","实际路径与本轮工具"),
        ("检查点","thread_id、messages、history_version","模型与工具执行状态")],[2.3,7.2,6.1])
    h("2.4　接口设计")
    p("普通接口统一返回code、message和data，成功code为0；错误HTTP状态与code一致。聊天请求保留教程的message和session_id，改名传title。接口清单见表4。健康检查不调用模型，即使尚未配置Key也可正常响应。")
    table("表4　主要接口",["方法与路径","职责"],[
        ("GET/POST /api/sessions","会话列表与新建"),("PATCH/DELETE /api/sessions/{id}","改名与删除"),
        ("GET /api/sessions/{id}/messages","历史问答"),("POST /api/chat、/api/agent/chat","普通与智能体"),
        ("POST /api/smart/chat、/api/chat/stream","自动路由与普通流式")],[9,6.6])
    p("422表示参数问题，404表示未知会话，409表示生成冲突，429表示限流，503表示模型未配置，502和504分别表示上游失败和超时。上游诊断仅记录异常类型，不把Key或错误原文交给前端。")

    document.add_page_break()
    h("3　核心功能实现",1);h("3.1　功能模块划分")
    p("后端config负责配置，models和database负责ORM与连接，memory负责会话和事务，llm负责模型与意图，tools负责工具，agent负责检查点，main提供路由，locks与rate_limit保证互斥和请求边界。前端App协调会话，Sidebar提供管理，ChatPanel执行聊天，stream模块解析增量事件。")
    h("3.2　关键功能实现");h("3.2.1　模型调用与结构化抽取",3)
    p("第一课演示脚本分别实现OpenAI SDK裸调用、LangChain普通对话、命令行流式、手动历史、最小FastAPI和结构化抽取。CourseInfo与BookInfo使用Pydantic Schema，BookInfo价格为float。真实能力验证抽取了课程教师和时间，书籍价格为45.0，结果保存于model-capabilities.json。导入脚本不会自动产生模型请求。")
    code("backend/app/schemas.py",'class BookInfo(BaseModel):\n    name: str\n    author: str\n    price: float  # 课后练习要求浮点价格\n    category: str')
    p("模型资源延迟初始化，应用生命周期关闭数据库、检查点及普通、分类、智能体模型的同步和异步HTTP连接池；启动失败仍清理已创建资源。环境变量优先于本地配置。模型名称、地址、超时和上下文窗口均可配置，依赖经过安装、pip check与Linux镜像构建后固定。")

    document.add_page_break()
    h("3.2.2　智能体工具调用",3)
    p("每个工具具有类型注解、参数和触发说明。query_order查询固定订单，get_weather查询少量城市，get_coupon提供优惠状态及条件。未知数据返回明确缺失，不自行编造；相应回答标注教学模拟。get_current_time使用Asia/Shanghai真实系统时钟，calculate执行确定性运算，两者不是模拟业务数据。")
    table("表5　五个工具与数据边界",["工具","输入与输出","边界"],[
        ("query_order","订单号→状态、物流","四个样例"),("get_weather","城市→天气、温度","未知城市无数据"),
        ("calculate","算术表达式→数值","AST与Decimal"),("get_current_time","无参数→北京时间","系统真实时钟"),
        ("get_coupon","优惠码→折扣条件","模拟有效、失效样例")],[4.3,6.3,5])
    p("calculate只接受数字、括号、正负号和四则运算，禁止函数、变量、幂和科学计数法；限制长度、AST节点、括号深度及数值规模，处理除零。AST会消除冗余括号，所以源码还要单独扫描括号深度。3874乘239结果925886由Decimal计算，不靠模型猜测。")
    code("backend/app/agent.py",'# model为缓存的模型；create_agent负责模型→工具→观察\nself._graph = create_agent(\n    model=model,\n    tools=TOOLS,\n    system_prompt=SYSTEM_PROMPT + AGENT_RULE,\n    checkpointer=self.saver,\n    state_schema=ChatAgentState,\n)\n# 每个业务会话对应稳定thread_id\nconfig = {"configurable": {"thread_id": session_id}}')
    p("智能体限制执行步数。通过本轮HumanMessage的唯一ID定位新增消息，再提取工具调用，避免把上轮工具重复展示到闲聊回复。")

    document.add_page_break()
    h("3.2.3　持久记忆与模式切换",3)
    p("业务历史是用户已完成问答的事实来源，检查点保存完整执行链。连续智能体调用版本一致时续跑；普通模式增加历史、失败或接近上下文上限时，下轮清理旧检查点，从业务历史重建。成功模型输出和业务事务完成后才确认本轮完成。双库没有跨文件原子事务，因此不能只依赖检查点宣称保存成功。")
    code("backend/app/agent.py",'# 检查点版本与业务历史不同，重建已完成上下文\nneed_rebuild = state.get("history_version", -1) != version\nif need_rebuild:\n    self.clear(session_id)  # 官方delete_thread\n    pending = to_messages(history)\nturn_id = uuid4().hex\npending.append(HumanMessage(content=message, id=turn_id))\nresult = self.graph.invoke(\n    {"messages": pending, "history_version": version + 1},\n    config,\n)')
    p("默认模型保留最近20个完整问答轮次，最多16000字符；完整历史仍保存。离线增长实验中，保存100轮共17780字符，实际上下文仅最近20轮3560字符，证明窗口限制生效；字符数不等于模型token数。重建舍弃过往工具中间过程，保留完成问答，本轮调用与结果由框架成对管理，不宣称无限记忆。")
    p("同一会话生成、改名与删除共享锁，重复生成或生成中删除返回409。用户问题与助手回答在同一事务中提交；失败不产生孤立问题，删除先清理检查点，错误时不向界面报告成功。真实集成中普通模式记录喜欢蓝色，智能体追问正确；服务重启后姓名和颜色仍正确恢复。")
    h("3.2.4　自动意图路由",3)
    p("Intent以Pydantic约束chat或agent，分类输入包括当前问题与必要最近历史。普通知识问答进入chat，订单和优惠券进入agent。前端自动模式实际调用统一入口，而不是在浏览器用关键词假装分类。分类失败返回友好错误，手动模式保留用于实验。")

    document.add_page_break()
    h("3.2.5　Vue交互与普通流式",3)
    p("Vue界面如图2所示。侧栏显示会话、改名与删除；消息区展示正文、本轮工具标签及实际路径。Enter发送，Shift加Enter换行，组合输入事件避免中文选字误发送。生成中禁止会话切换和重复发送，失败保留问题供重试。刷新按本地会话ID恢复，已删除会话回到可用记录。")
    picture(ROOT/"artifacts/evidence/frontend-agent.png","图2　自动路由与订单工具展示")
    p("普通模式使用fetch接收POST SSE。TextDecoder以stream方式解码中文，网络块先进入缓冲，遇到完整空行事件才解析JSON。接口状态、error事件和DONE都校验；自动与智能体模式本轮为完整响应。工具订单、天气及优惠券为模拟数据，界面始终提示边界。")
    p("后端完整生成并保存成功后才发送DONE；断开或生成失败不保存部分答案。TCP断流测试实际关闭客户端连接，验证历史仍为空、生成时删除409、取消后可改名。极端情况下网络在提交后中断，客户端无法确认是否已保存，所以重试前需刷新记录。")

    document.add_page_break()
    h("4　测试与优化",1);h("4.1　测试方案与用例")
    p("测试分为离线回归、真实模型、浏览器和容器。离线测试注入确定性模型，真实执行LangGraph循环；真实评估使用独立检查点目录，每题唯一会话。CI不配置真实Key，但GitHub账号账单问题导致任务未启动，不能记作云端通过。本机分层验证与当次结果见表6。")
    table("表6　分层验证结果",["层次","主要内容","当次结果"],[
        ("后端离线","接口、恢复、断流、限流、资源释放","32通过"),("前端离线","UTF-8、跨块JSON、CRLF、错误","4通过"),
        ("构建依赖","pip check、npm build、Docker build","通过"),("模型能力","普通、CourseInfo、BookInfo、流式","4通过"),
        ("真实API","多轮、模式、路由、SSE、重启","9通过"),("浏览器","管理、模式、标签、刷新、删除","通过"),
        ("容器","代理、持久化、备份、限流","13通过")],[3,9,3.6])
    p("未知会话测试确认404且不调用模型；空白和超长输入422；模型错误502且不泄露异常中的secret；并发生成与删除409。工具测试覆盖表达式注入、幂、除零、极大数和未知业务数据。流式测试同时覆盖成功提交、上游失败、真实连接取消。")
    h("4.2　测试结果分析")
    p("2026年10月5日完成真实验证。课堂固定10题两轮均10/10，包含两题订单、两题天气、两题计算、两题时间和两题无需工具的问题。补充优惠券、多工具、未知订单和未知城市用于检验第五工具与边界。每次输出JSON及CSV，记录期望、实际、耗时和错误。")

    document.add_page_break()
    h("4.2.1　工具选择评估",3)
    p("固定题评估如表7所示。课堂指标要求期望工具包含于实际集合，无工具题要求实际为空；同时保存集合完全相等的更严格指标。首轮课堂10/10，补充3/4，NOTFOUND订单未触发工具。加强未知订单触发说明后第二轮课堂10/10、补充4/4。未删除首轮结果。")
    table("表7　固定工具选择评估",["场景","期望工具","优化后"],[
        ("订单两题","query_order","2/2"),("天气两题","get_weather","2/2"),("计算两题","calculate","2/2"),
        ("时间两题","get_current_time","2/2"),("闲聊两题","无工具","2/2"),("优惠券","get_coupon","1/1"),
        ("时间和一年分钟数","time + calculate","1/1"),("未知订单和城市","order / weather","2/2")],[6.2,6.4,3])
    p("补充条件任务先查询订单，已发货才查询当前时间，工具按条件完成。动态未来日期题使用当前日期加30天，避免固定过去日期失效。这些补充题同时观察最终正文，不能只把工具标签当成成功。")
    h("4.2.2　提示词对照实验",3)
    p("两个独立会话分别保留和去掉强制工具规则，询问同一个订单。本次两组均调用query_order并正确返回样例状态，未观察到工具选择差异。工具描述和基础系统提示仍保留，因此实验只能说明这次样例没有变化，不能推出强制规则永远无效。结果保存在prompt-experiment.json。")

    document.add_page_break()
    h("4.3　问题与优化")
    p("日期题初测虽然调用时间和计算工具，却引用了训练日期2024年5月20日，得到619天。该记录最初只按工具选择标记，不能算答案通过。增加必须依据本轮时钟的规则并核对答案后，最终使用2026年10月5日到11月4日，得到30天。回答还出现中间自我校正，说明最终数值正确不意味着解释过程完全稳定。")
    table("表8　实际问题及修正",["问题","处理","验证"],[
        ("未知订单未触发工具","完善说明与规则","补充4题全通过"),("日期用了训练数据","本轮时钟约束、答案断言","最终30天"),
        ("AST忽略冗余括号","源码扫描深度","边界测试通过"),("925,886被误判错误","校验归一化千位逗号","Docker重测通过"),
        ("8080被其他项目占用","端口参数化用8081","不干扰现有服务"),("前端依赖Node版本不匹配","固定Element Plus2.11.8","构建与测试通过")],[4.5,6.5,4.6])
    p("前端按需注册组件将JS构建产物降至约252KB。模型异常不直接暴露给浏览器，日志记录路径、状态、耗时和异常类型。评估保留失败、优化与重测过程，避免只展示最好一次。")
    h("4.4　可靠性与局限")
    p("模型输出存在随机性，temperature为0也不能保证永远一致。固定样本100%只代表该轮小样本工具选择。系统未提供真实订单与天气、登录鉴权、RAG、多实例共享限流或公网交付；上下文受窗口限制，双库恢复依赖版本与清理策略。当前工程验证覆盖计划核心范围，后续扩展需新测试。")

    document.add_page_break()
    h("5　部署方案",1);h("5.1　部署架构")
    p("后端基于Python3.11-slim镜像，使用固定依赖，仅复制app与requirements。前端以Node22多阶段npm ci构建，最终交给Nginx，不携带node_modules。Compose注入backend/.env，挂载双库目录，设置Asia/Shanghai，后端监听容器0.0.0.0:8000但不映射宿主机端口。")
    p("前端默认监听宿主机127.0.0.1:8080，本机已有其他Docker应用占用，所以实际配置8081，保留默认端口供其他机器复现。Nginx覆盖X-Forwarded-For为真实连接地址，后端仅信任172.29.90.0/24代理网段；关闭SSE缓冲，读取超时180秒。")
    h("5.2　部署步骤")
    code("项目根目录 PowerShell",'# 确认Docker Desktop Linux引擎和本地模型配置\ndocker compose up -d --build --wait\ndocker compose ps\n# 查看日志，不输出.env内容\ndocker compose logs --tail 50 backend\n# 停止本项目，挂载数据仍保留\ndocker compose down')
    p("后端健康检查调用不触发模型的/api/health，前端等待后端健康再启动。依赖目录、真实.env与数据库均通过Git忽略和构建复制范围排除；模型配置只在后端进程，浏览器无法获取Key。源码包包含无密钥样例和运行说明。")
    h("5.3　运行限制")
    p("应用为单worker，内存滑动窗口每IP60秒20次，重启清零，健康检查豁免。不宣称多实例共享限流。开发后端和Docker后端不要同时写同一数据目录。改变代理子网必须同步修改Uvicorn可信地址再构建。")

    document.add_page_break()
    h("5.4　持久化与恢复")
    p("业务库smartchat.db与检查点库agent_memory.db在同一宿主机挂载目录。验收先完成普通、智能体、自动及流式，记录历史；down再up、重建镜像并force-recreate后，历史内容均与原记录一致。镜像与运行数据分离保证重建程序不删除问答。")
    table("表9　Docker验收结果",["项目","结果"],[
        ("页面与健康、三模式、SSE","通过"),("停启后历史相同","通过"),("重建镜像和容器后历史相同","通过"),
        ("SQLite备份完整性","通过"),("独立目录恢复历史、继续智能体","通过"),
        ("伪造转发头与第21次请求429","通过"),("Retry-After与健康豁免","通过")],[12,3.6])
    p("备份必须等生成结束并停止后端，以免两库来自不同轮次。backup_data.py使用SQLite backup API复制并检查integrity_check，拒绝覆盖已有目标。恢复测试把两库复制到独立临时目录，启动新FastAPI，确认历史一致并真实追问姓名，回答小明；避免直接覆盖当前演示数据。")
    code("备份步骤 PowerShell",'# 先停止后端，取得双库完成状态\ndocker compose stop backend\n.venv\\Scripts\\python.exe backend\\backup_data.py artifacts/private/backup-new\n# 保留备份后恢复运行\ndocker compose up -d --wait')
    p("正式恢复先保留原data目录，再建立空目录复制两份备份，不能混入旧WAL和SHM。部署证据与全部13项验收结果在docker-acceptance.json，个人数据库和备份不上传公开仓库。")

    document.add_page_break()
    h("6　个人总结与展望",1);h("6.1　项目完成情况")
    p("本项目完成三课核心功能和所选课后练习：BookInfo抽取、优惠券工具、提示词对照、条件多工具、自动评估与Docker交付。GitHub通过阶段分支与PR记录实现，发布标签在对应验收后创建。报告、截图、API结果及代码应保持一致，不预填未来答辩成绩或教师确认。")
    p("报告与代码整理采用AI编程辅助，个人需逐模块理解、复核和修改叙述，并亲自完成演示与答辩。本文不把自动化测试等同于本人已掌握，也不虚构课堂开发天数。最终提交前还需核对原始学校DOCX模板、更新目录与页码、进行本人演练；答辩后经老师确认再提交。")
    h("6.2　工程改进与难点")
    p("工程改进包括第五优惠券工具、安全算术、完整轮次事务、仅本轮工具记录、双库版本恢复、生成锁、中文SSE缓冲与断流取消、代理可信地址以及备份恢复验证。主要难点是模型与业务状态的协调，以及工具正确调用仍可能产生错误最终答案。通过版本与失败清理提高可恢复性，通过答案检查补充工具指标。")
    h("6.3　后续展望")
    p("后续可接入真实天气和订单、登录鉴权、RAG知识库及公网部署，增加智能体逐Token过程与多实例共享限流。应先明确数据权限与业务需求，再做新的集成和验收。本轮专注本机稳定演示和课程要求，不把扩展计划当作已实现功能。")
    h("6.4　参考资料")
    for text in [
        "[1] 课程教学资料. 第1—3课LangChain与智能体开发笔记，2026.",
        "[2] 太原理工大学23级创新训练营学生版要求[EB/OL]. https://www.kdocs.cn/l/ccnOFR0hHXak，访问日期2026-10-05.",
        "[3] LangChain. Agents[EB/OL]. https://docs.langchain.com/oss/python/langchain/agents，访问日期2026-10-05.",
        "[4] 阿里云. 百炼OpenAI兼容接口[EB/OL]. https://help.aliyun.com/zh/model-studio/qwen-api-via-openai-chat-completions，访问日期2026-10-05.",
    ]:
        para=p(text);para.paragraph_format.first_line_indent=Pt(0)
        para.paragraph_format.line_spacing=1.0
        for r in para.runs:font(r,size=10.5)
    output.parent.mkdir(parents=True,exist_ok=True)
    document.save(output)
    print("已生成报告初稿；原始学校DOCX模板核对和本人复核仍待完成。")


if __name__ == "__main__":
    parser=argparse.ArgumentParser()
    parser.add_argument("--profile",type=Path,required=True)
    parser.add_argument("--output",type=Path,required=True)
    args=parser.parse_args()
    args.output.parent.mkdir(parents=True,exist_ok=True)
    build(args.profile,args.output)
